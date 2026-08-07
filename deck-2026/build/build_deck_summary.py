import sys
sys.path.insert(0, "/root/.claude/skills/brand-design/scripts")
import brand
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor

OUT = "/tmp/claude-0/-home-user-Claude-Code-First/5fca47e4-c527-52de-b9be-2f51cb4c5b55/scratchpad/dataminr_2026_alerting_summary.pptx"

prs = brand.new_deck()

MARGIN = 0.7
GAP = 0.35
COL_W = (13.333 - 2 * MARGIN - GAP) / 2
COL_X = [MARGIN, MARGIN + COL_W + GAP]


def fmt_k(v):
    if v >= 1000:
        return f"{v/1000:.0f}k" if v % 1000 else f"{int(v/1000)}k"
    return str(int(v))


def divider_line(slide, y0, y1):
    x = MARGIN + COL_W + GAP / 2
    ln = slide.shapes.add_connector(1, Inches(x), Inches(y0), Inches(x), Inches(y1))
    ln.line.color.rgb = RGBColor.from_string(brand.LIGHT_GRAY)
    ln.line.width = Pt(0.75)
    brand.flatten(ln)


def column_header(slide, x, y, w, text):
    tx = brand._textbox(slide, x, y, w, 0.32)
    brand._run(tx.text_frame.paragraphs[0], text, font=brand.FONT_HEADLINE, size=16, color=brand.SENTINEL_BLUE)
    return y + 0.32


def hero_block(slide, x, y, w, eyebrow, number, descriptor, caveat=None, number_size=76,
               eyebrow_gap=0.24, desc_h=0.32, caveat_h=0.3):
    tx = brand._textbox(slide, x, y, w, 0.22)
    brand._run(tx.text_frame.paragraphs[0], eyebrow, size=10.5, color=brand.MEDIUM_GRAY)
    num_h = number_size / 72.0 * 1.05
    y2 = y + eyebrow_gap
    tx = brand._textbox(slide, x, y2, w, num_h)
    brand._run(tx.text_frame.paragraphs[0], number, font=brand.FONT_HEADLINE, size=number_size, color=brand.CIPHER_BLUE)
    y3 = y2 + num_h
    tx = brand._textbox(slide, x, y3, w, desc_h)
    brand._run(tx.text_frame.paragraphs[0], descriptor, size=12.5, color=brand.PLATFORM_GRAY)
    y4 = y3 + desc_h
    if caveat:
        tx = brand._textbox(slide, x, y4, w, caveat_h)
        brand._run(tx.text_frame.paragraphs[0], caveat, size=9.5, italic=True, color=brand.MEDIUM_GRAY)
        y4 += caveat_h
    return y4


def indicator_row(slide, x, y, w, h, icon_name, title, desc):
    brand.rounded_box(slide, Inches(x), Inches(y), Inches(w), Inches(h), brand.ANALYTIC_GRAY)
    icon_h = 0.26
    slide.shapes.add_picture(brand.icon(icon_name, "line"), Inches(x + 0.14), Inches(y + (h - icon_h) / 2), height=Inches(icon_h))
    tx = brand._textbox(slide, x + 0.5, y + 0.06, w - 0.62, h - 0.12, anchor=MSO_ANCHOR.MIDDLE)
    p = tx.text_frame.paragraphs[0]
    brand._run(p, title + " — ", size=11, bold=True, color=brand.SENTINEL_BLUE)
    brand._run(p, desc, size=10.5, color=brand.PLATFORM_GRAY)


def bar_chart(slide, x, y, w, h, data, y_max, *, label_every=1, note_partial_idx=None):
    brand.rounded_box(slide, Inches(x), Inches(y), Inches(w), Inches(h), brand.ANALYTIC_GRAY)
    pad_x, pad_top, pad_bottom = 0.4, 0.2, 0.42
    plot_x = x + pad_x
    plot_w = w - 2 * pad_x
    plot_y = y + pad_top
    plot_h = h - pad_top - pad_bottom
    n = len(data)
    gap_frac = 0.35
    slot_w = plot_w / n
    bar_w = slot_w * (1 - gap_frac)

    for frac in (0.0, 0.5, 1.0):
        gy = plot_y + plot_h * (1 - frac)
        ln = slide.shapes.add_connector(1, Inches(plot_x), Inches(gy), Inches(plot_x + plot_w), Inches(gy))
        ln.line.color.rgb = RGBColor.from_string(brand.LIGHT_GRAY)
        ln.line.width = Pt(0.5)
        brand.flatten(ln)
        lbl = brand._textbox(slide, x, gy - 0.08, pad_x - 0.06, 0.16, anchor=MSO_ANCHOR.MIDDLE)
        p = lbl.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.RIGHT
        brand._run(p, fmt_k(y_max * frac), size=7, color=brand.MEDIUM_GRAY)

    for i, (label, val) in enumerate(data):
        bx = plot_x + i * slot_w + (slot_w - bar_w) / 2
        is_partial = note_partial_idx is not None and i == note_partial_idx
        if val is None:
            bh = 0.03
            by = plot_y + plot_h - bh
            box = brand.rounded_box(slide, Inches(bx), Inches(by), Inches(bar_w), Inches(bh), brand.CIPHER_BLUE, radius=0.5)
            box.fill.fore_color.rgb = RGBColor.from_string(brand.LIGHT_GRAY)
        else:
            frac = max(val / y_max, 0.008)
            bh = plot_h * frac
            by = plot_y + plot_h - bh
            box = brand.rounded_box(slide, Inches(bx), Inches(by), Inches(bar_w), Inches(bh), brand.CIPHER_BLUE, radius=0.35)
            if is_partial:
                box.fill.fore_color.rgb = RGBColor.from_string(brand.PERCEPTION_ALLOY)
        if i % label_every == 0:
            lx = plot_x + i * slot_w
            ltx = brand._textbox(slide, lx, plot_y + plot_h + 0.04, slot_w, 0.3)
            p = ltx.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            brand._run(p, label, size=6.5, color=brand.MEDIUM_GRAY)


# ============================================================ slide 1 — current trends

s = brand.add_content_slide(
    prs, title="Current alerting trends",
    subtitle="Realized alert volume across two targeted expansions",
    background="light", page_no=1,
    sources=["Internal alert volume export, Vessels of Interest / sanctioned-vessel theme",
             "Internal alert volume export, Strait of Hormuz chokepoint theme"],
)
divider_line(s, 1.85, 6.5)

# -- left: Vessels of interest
x = COL_X[0]
y = column_header(s, x, 1.85, COL_W, "Vessels of interest & sanctions")
y = hero_block(s, x, y + 0.08, COL_W, "GROWTH, MAY 2025 -> MAR 2026", "~59x",
               "13,351 alerts at peak (Mar 2026), up from 228")
voi_data = [
    ("Mar '25", None), ("Apr '25", None), ("May '25", 228), ("Jun '25", 1209), ("Jul '25", 3179),
    ("Aug '25", 4436), ("Sep '25", 3674), ("Oct '25", 7713), ("Nov '25", 11740), ("Dec '25", 12470),
    ("Jan '26", 12452), ("Feb '26", 12389), ("Mar '26", 13351), ("Apr '26", 9581), ("May '26", 11505), ("Jun '26", 8779),
]
bar_chart(s, x, y + 0.08, COL_W, 6.5 - (y + 0.08), voi_data, y_max=14000, label_every=4)

# -- right: Strait of Hormuz
x = COL_X[1]
y = column_header(s, x, 1.85, COL_W, "Strait of Hormuz & MENA expansion")
y = hero_block(s, x, y + 0.08, COL_W, "PEAK MONTHLY ALERTS", "903",
               "Up from zero in Feb 2026 -- stood up for Epic Fury")
hormuz_data = [
    ("Feb '26", 0), ("Mar '26", 0), ("Apr '26", 132), ("May '26", 769), ("Jun '26", 330), ("Jul '26", 903), ("Aug '26", 80),
]
bar_chart(s, x, y + 0.08, COL_W, 6.5 - (y + 0.08), hormuz_data, y_max=900, label_every=1, note_partial_idx=6)

# ============================================================ slide 2 — upcoming releases

s = brand.add_content_slide(
    prs, title="Upcoming releases & projections",
    subtitle="Test-set projections for themes launching next week",
    background="light", page_no=2,
    sources=["Internal test-set projections, MAV alerting theme (pre-launch)",
             "Internal test-set projections, Arctic alerting theme (pre-launch)"],
)
divider_line(s, 1.85, 6.5)

ROW_H = 0.6
ROW_GAP = 0.08

# -- left: MAV
x = COL_X[0]
y = column_header(s, x, 1.85, COL_W, "South China Sea -- MAV alerting")
y = hero_block(s, x, y + 0.06, COL_W, "PROJECTED DAILY VOLUME", "~300", "alerts/day at launch",
               caveat="Test-set estimate -- launching week of Aug 10, 2026",
               number_size=56, eyebrow_gap=0.22, desc_h=0.26, caveat_h=0.22)
rows = [
    ("signal", "Dark activity", "AIS/transponder signal loss consistent with evasive behavior."),
    ("navigation", "Military zone entry", "Course tracks into designated military or restricted zones."),
    ("network", "Vessel meeting", "Close-proximity rendezvous between vessels at sea."),
]
ry = y + 0.1
for icon_name, title, desc in rows:
    indicator_row(s, x, ry, COL_W, ROW_H, icon_name, title, desc)
    ry += ROW_H + ROW_GAP
tx = brand._textbox(s, x, ry + 0.02, COL_W, 0.22)
brand._run(tx.text_frame.paragraphs[0], "Per-indicator daily split not yet available from the test set.",
           size=9, italic=True, color=brand.MEDIUM_GRAY)

# -- right: Arctic
x = COL_X[1]
y = column_header(s, x, 1.85, COL_W, "Arctic expansion")
y = hero_block(s, x, y + 0.06, COL_W, "PROJECTED DAILY VOLUME", "~35", "alerts/day at launch",
               caveat="Test-set estimate -- launching week of Aug 10, 2026",
               number_size=56, eyebrow_gap=0.22, desc_h=0.26, caveat_h=0.22)
rows = [
    ("navigation", "EEZ entries", "Vessel entries into exclusive economic zones of Arctic states."),
    ("signal", "Dark activity", "AIS/transponder signal loss for medium- and high-risk vessels."),
]
ry = y + 0.1
for icon_name, title, desc in rows:
    indicator_row(s, x, ry, COL_W, ROW_H, icon_name, title, desc)
    ry += ROW_H + ROW_GAP
tx = brand._textbox(s, x, ry + 0.02, COL_W, 0.22)
brand._run(tx.text_frame.paragraphs[0], "Split by risk tier (medium vs. high) not yet available from the test set.",
           size=9, italic=True, color=brand.MEDIUM_GRAY)

prs.save(OUT)
print("saved", OUT)
