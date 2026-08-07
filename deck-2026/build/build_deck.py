import sys
sys.path.insert(0, "/root/.claude/skills/brand-design/scripts")
import brand
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor

OUT = "/tmp/claude-0/-home-user-Claude-Code-First/5fca47e4-c527-52de-b9be-2f51cb4c5b55/scratchpad/dataminr_2026_alerting_expansion.pptx"

prs = brand.new_deck()

# ---------------------------------------------------------------- helpers

def stat_row(slide, stats, y=1.85, h=1.05):
    """stats: list of (label, value, unit) tuples, drawn as equal-width Analytic Gray tiles."""
    n = len(stats)
    margin = 0.7
    gap = 0.2
    total_w = 13.333 - 2 * margin
    w = (total_w - gap * (n - 1)) / n
    for i, (label, value, unit) in enumerate(stats):
        x = margin + i * (w + gap)
        brand.rounded_box(slide, Inches(x), Inches(y), Inches(w), Inches(h), brand.ANALYTIC_GRAY)
        tx = brand._textbox(slide, x + 0.18, y + 0.13, w - 0.36, 0.3)
        brand._run(tx.text_frame.paragraphs[0], label, size=11, color=brand.MEDIUM_GRAY)
        tx = brand._textbox(slide, x + 0.18, y + 0.42, w - 0.36, 0.55)
        p = tx.text_frame.paragraphs[0]
        brand._run(p, value, font=brand.FONT_HEADLINE, size=26, color=brand.SENTINEL_BLUE)
        if unit:
            brand._run(p, " " + unit, size=13, color=brand.PLATFORM_GRAY)


def bar_chart(slide, x, y, w, h, data, y_max, *, label_every=1, note_partial_idx=None, note_zero_ok=True):
    """data: list of (label, value_or_None). None => no-data tick. Draws Cipher Blue bars in a card."""
    brand.rounded_box(slide, Inches(x), Inches(y), Inches(w), Inches(h), brand.ANALYTIC_GRAY)
    pad_x, pad_top, pad_bottom = 0.35, 0.35, 0.55
    plot_x = x + pad_x
    plot_w = w - 2 * pad_x
    plot_y = y + pad_top
    plot_h = h - pad_top - pad_bottom
    n = len(data)
    gap_frac = 0.35
    slot_w = plot_w / n
    bar_w = slot_w * (1 - gap_frac)

    # gridlines at 0/25/50/75/100%
    for frac in (0.0, 0.25, 0.5, 0.75, 1.0):
        gy = plot_y + plot_h * (1 - frac)
        ln = slide.shapes.add_connector(1, Inches(plot_x), Inches(gy), Inches(plot_x + plot_w), Inches(gy))
        ln.line.color.rgb = RGBColor.from_string(brand.LIGHT_GRAY)
        ln.line.width = Pt(0.5)
        brand.flatten(ln)
        lbl = brand._textbox(slide, x, gy - 0.09, pad_x - 0.06, 0.18, anchor=MSO_ANCHOR.MIDDLE)
        p = lbl.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.RIGHT
        brand._run(p, f"{int(y_max * frac):,}", size=7, color=brand.MEDIUM_GRAY)

    for i, (label, val) in enumerate(data):
        bx = plot_x + i * slot_w + (slot_w - bar_w) / 2
        is_partial = note_partial_idx is not None and i == note_partial_idx
        if val is None:
            bh = 0.04
            by = plot_y + plot_h - bh
            box = brand.rounded_box(slide, Inches(bx), Inches(by), Inches(bar_w), Inches(bh), brand.CIPHER_BLUE, radius=0.5)
            box.fill.fore_color.rgb = RGBColor.from_string(brand.LIGHT_GRAY)
        else:
            frac = max(val / y_max, 0.006)
            bh = plot_h * frac
            by = plot_y + plot_h - bh
            fill_hex = brand.CIPHER_BLUE
            box = brand.rounded_box(slide, Inches(bx), Inches(by), Inches(bar_w), Inches(bh), fill_hex, radius=0.4)
            if is_partial:
                box.fill.fore_color.rgb = RGBColor.from_string(brand.PERCEPTION_ALLOY)
        if i % label_every == 0:
            lx = plot_x + i * slot_w
            ltx = brand._textbox(slide, lx, plot_y + plot_h + 0.06, slot_w, 0.4)
            p = ltx.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            brand._run(p, label, size=8, color=brand.MEDIUM_GRAY)
            if is_partial:
                p2 = ltx.text_frame.add_paragraph()
                p2.alignment = PP_ALIGN.CENTER
                brand._run(p2, "(partial)", size=6.5, italic=True, color=brand.MEDIUM_GRAY)


def indicator_card(slide, x, y, w, h, icon_name, title, desc):
    brand.rounded_box(slide, Inches(x), Inches(y), Inches(w), Inches(h), brand.ANALYTIC_GRAY)
    icon_h = 0.32
    slide.shapes.add_picture(brand.icon(icon_name, "line"), Inches(x + 0.25), Inches(y + 0.16), height=Inches(icon_h))
    tx = brand._textbox(slide, x + 0.25, y + 0.16 + icon_h + 0.06, w - 0.5, 0.28)
    brand._run(tx.text_frame.paragraphs[0], title, font=brand.FONT_HEADLINE, size=14, color=brand.SENTINEL_BLUE)
    desc_y = y + 0.16 + icon_h + 0.06 + 0.30
    desc_h = max(h - (desc_y - y) - 0.12, 0.2)
    tx = brand._textbox(slide, x + 0.25, desc_y, w - 0.5, desc_h)
    brand._run(tx.text_frame.paragraphs[0], desc, size=10.5, color=brand.PLATFORM_GRAY)


def big_stat_slide(slide, eyebrow, number, descriptor, subline):
    tx = brand._textbox(slide, 0.7, 1.85, 6.0, 0.4)
    brand._run(tx.text_frame.paragraphs[0], eyebrow, font=brand.FONT_HEADLINE, size=16, color=brand.SENTINEL_BLUE)
    tx = brand._textbox(slide, 0.7, 2.15, 6.5, 1.9, anchor=MSO_ANCHOR.MIDDLE)
    brand._run(tx.text_frame.paragraphs[0], number, font=brand.FONT_HEADLINE, size=110, color=brand.CIPHER_BLUE)
    tx = brand._textbox(slide, 0.7, 4.05, 6.3, 0.5)
    brand._run(tx.text_frame.paragraphs[0], descriptor, size=18, color=brand.PLATFORM_GRAY)
    tx = brand._textbox(slide, 0.7, 4.5, 6.3, 0.6)
    brand._run(tx.text_frame.paragraphs[0], subline, size=12.5, italic=True, color=brand.MEDIUM_GRAY)


# ---------------------------------------------------------------- 1. title slide

brand.insert_title_slide(
    prs,
    title="2026 alerting expansion",
    subtitle="Targeted growth in strategic chokepoints and conflict response",
)

# ---------------------------------------------------------------- 2. framing divider

brand.add_section_divider(
    prs,
    "Targeted, not global",
    subtitle="Alert growth in 2026 came from focused expansion, not a global volume increase",
    page_no=2,
)

# ---------------------------------------------------------------- 3. agenda

s = brand.add_content_slide(prs, title="Four places we expanded", background="light", page_no=3,
                             sources=["Internal alert volume export, Vessels of Interest / sanctioned-vessel theme",
                                      "Internal alert volume export, Strait of Hormuz chokepoint theme",
                                      "Internal test-set projections, MAV alerting theme (pre-launch)",
                                      "Internal test-set projections, Arctic alerting theme (pre-launch)"])
tx = brand._textbox(slide=s, x=0.7, y=1.7, w=11.93, h=4.8)
tf = tx.text_frame
tf.word_wrap = True
brand.add_bullet_paragraph(tf, [
    {"text": "Vessels of interest & sanctions", "bold": True, "color": brand.SENTINEL_BLUE},
    {"text": " — Russian shadow-fleet designations drove alert volume up ~59x since May 2025.", "color": brand.PLATFORM_GRAY},
], first=True)
brand.add_bullet_paragraph(tf, [
    {"text": "Strait of Hormuz & MENA expansion", "bold": True, "color": brand.SENTINEL_BLUE},
    {"text": " — new chokepoint alerting stood up in direct response to Epic Fury: 0 to 900+ alerts/month in 4 months.", "color": brand.PLATFORM_GRAY},
])
brand.add_bullet_paragraph(tf, [
    {"text": "South China Sea — MAV alerting", "bold": True, "color": brand.SENTINEL_BLUE},
    {"text": " — launching next week: ~300 projected alerts/day across three behavioral risk indicators.", "color": brand.PLATFORM_GRAY},
])
brand.add_bullet_paragraph(tf, [
    {"text": "Arctic expansion", "bold": True, "color": brand.SENTINEL_BLUE},
    {"text": " — launching next week: ~35 projected alerts/day across EEZ entries and dark activity.", "color": brand.PLATFORM_GRAY},
])
tx2 = brand._textbox(s, 0.7, 6.05, 11.93, 0.5)
brand._run(tx2.text_frame.paragraphs[0],
           "A separate project to generalize and globalize these alert themes exists but has not started.",
           size=12.5, italic=True, color=brand.MEDIUM_GRAY)

# ---------------------------------------------------------------- 4. vessels of interest

s = brand.add_content_slide(
    prs, title="Vessels of interest & sanctions",
    subtitle="Russian shadow-fleet designations drove a sustained alert plateau",
    background="light", page_no=4,
    sources=["Internal alert volume export, Vessels of Interest / sanctioned-vessel theme"],
)
stat_row(s, [("May 2025 (first month)", "228", "alerts"), ("Peak month (Mar 2026)", "13,351", "alerts"), ("Growth, May '25 -> Mar '26", "~59x", "")])
voi_data = [
    ("Mar '25", None), ("Apr '25", None), ("May '25", 228), ("Jun '25", 1209), ("Jul '25", 3179),
    ("Aug '25", 4436), ("Sep '25", 3674), ("Oct '25", 7713), ("Nov '25", 11740), ("Dec '25", 12470),
    ("Jan '26", 12452), ("Feb '26", 12389), ("Mar '26", 13351), ("Apr '26", 9581), ("May '26", 11505), ("Jun '26", 8779),
]
bar_chart(s, 0.7, 3.15, 11.93, 3.15, voi_data, y_max=14000, label_every=2)

# ---------------------------------------------------------------- 5. strait of hormuz

s = brand.add_content_slide(
    prs, title="Strait of Hormuz & MENA expansion",
    subtitle="Stood up in direct response to Epic Fury",
    background="light", page_no=5,
    sources=["Internal alert volume export, Strait of Hormuz chokepoint theme"],
)
stat_row(s, [("Feb-Mar 2026 (zero coverage)", "0", "alerts"), ("Peak month (Jul 2026)", "903", "alerts"), ("Time to peak coverage", "4", "months")])
hormuz_data = [
    ("Feb '26", 0), ("Mar '26", 0), ("Apr '26", 132), ("May '26", 769), ("Jun '26", 330), ("Jul '26", 903), ("Aug '26", 80),
]
bar_chart(s, 0.7, 3.15, 11.93, 3.15, hormuz_data, y_max=900, label_every=1, note_partial_idx=6)

# ---------------------------------------------------------------- 6. MAV (projected)

s = brand.add_content_slide(
    prs, title="South China Sea -- MAV alerting",
    subtitle="Launching next week -- projected volume from test-set output",
    background="light", page_no=6,
    sources=["Internal test-set projections, MAV alerting theme (pre-launch)"],
)
big_stat_slide(s, "PROJECTED DAILY VOLUME", "~300", "alerts/day at launch",
               "Test-set estimate, not a realized alert count -- launching week of Aug 10, 2026")
indicator_card(s, 7.55, 2.0, 5.08, 1.4, "signal", "Dark activity", "AIS/transponder signal loss consistent with evasive behavior.")
indicator_card(s, 7.55, 3.55, 5.08, 1.4, "navigation", "Military zone entry", "Course tracks into designated military or restricted zones.")
indicator_card(s, 7.55, 5.10, 5.08, 1.4, "network", "Vessel meeting", "Close-proximity rendezvous between vessels at sea.")
tx = brand._textbox(s, 0.7, 5.4, 6.3, 0.9)
brand._run(tx.text_frame.paragraphs[0],
           "Per-indicator daily split not yet available from the test set.",
           size=11.5, italic=True, color=brand.MEDIUM_GRAY)

# ---------------------------------------------------------------- 7. arctic (projected)

s = brand.add_content_slide(
    prs, title="Arctic expansion",
    subtitle="Launching next week -- projected volume from test-set output",
    background="light", page_no=7,
    sources=["Internal test-set projections, Arctic alerting theme (pre-launch)"],
)
big_stat_slide(s, "PROJECTED DAILY VOLUME", "~35", "alerts/day at launch",
               "Test-set estimate, not a realized alert count -- launching week of Aug 10, 2026")
indicator_card(s, 7.55, 2.6, 5.08, 1.55, "navigation", "EEZ entries", "Vessel entries into exclusive economic zones of Arctic states.")
indicator_card(s, 7.55, 4.35, 5.08, 1.55, "signal", "Dark activity", "AIS/transponder signal loss for medium- and high-risk vessels.")
tx = brand._textbox(s, 0.7, 5.4, 6.3, 0.9)
brand._run(tx.text_frame.paragraphs[0],
           "Split by risk tier (medium vs. high) not yet available from the test set.",
           size=11.5, italic=True, color=brand.MEDIUM_GRAY)

# ---------------------------------------------------------------- 8. what's next

s = brand.add_content_slide(prs, title="What's next", background="dark", page_no=8)
tx = brand._textbox(s, 0.7, 1.9, 11.93, 4.6)
tf = tx.text_frame
tf.word_wrap = True
brand.add_bullet_paragraph(tf, [
    {"text": "Launch MAV and Arctic alerting", "bold": True, "color": brand.WHITE},
    {"text": " and replace projected volume with realized coverage.", "color": brand.WHITE},
], first=True)
brand.add_bullet_paragraph(tf, [
    {"text": "Cross-check Vessels of Interest/Hormuz volume", "bold": True, "color": brand.WHITE},
    {"text": " against OFAC/EU designation dates and Epic Fury escalation timing.", "color": brand.WHITE},
])
brand.add_bullet_paragraph(tf, [
    {"text": "Fold in Windward-sourced leads", "bold": True, "color": brand.WHITE},
    {"text": " as a validation case study once specific examples are confirmed.", "color": brand.WHITE},
])
brand.add_bullet_paragraph(tf, [
    {"text": "Scope the generalize-and-globalize project", "bold": True, "color": brand.WHITE},
    {"text": " for existing alert themes -- not yet started.", "color": brand.WHITE},
])

prs.save(OUT)
print("saved", OUT)
