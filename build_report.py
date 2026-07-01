"""Build a three-tab Excel report: predictions vs GT, performance, and images."""

import csv
import io
import os
import re
import ssl
import urllib.request
from collections import defaultdict
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

PREDICTIONS_PATH = "/root/.claude/uploads/4198d3e6-bbc8-5bb5-b7be-6a36a4116264/96b7eeb3-predictions.csv"
TESTSHEET_PATH   = "/root/.claude/uploads/4198d3e6-bbc8-5bb5-b7be-6a36a4116264/2c51133a-testsheet.csv"
CORRECTIONS_PATH = "/root/.claude/uploads/4198d3e6-bbc8-5bb5-b7be-6a36a4116264/cd45af46-Untitled_spreadsheet__false_positives_1.csv"
OUT_PATH         = "headline_report.xlsx"
IMG_CACHE_DIR    = "/tmp/claude-0/-home-user-Claude-Code-First/4198d3e6-bbc8-5bb5-b7be-6a36a4116264/scratchpad/img_cache"
IMG_W, IMG_H     = 200, 150   # display size in pixels

GT_FIELDS    = ["police_1","police_2plus","fire_truck_1","fire_truck_2plus",
                "ambulance_1","ambulance_2plus","no_emergency_vehicles","hard_to_tell",
                "crash","pulled_over","blocked_road","construction","fire","crowd",
                "other_identifiable","no_incident_visible"]
GT_COL_START = 3

GREEN  = PatternFill("solid", fgColor="C6EFCE")
RED    = PatternFill("solid", fgColor="FFC7CE")
YELLOW = PatternFill("solid", fgColor="FFEB9C")
GREY   = PatternFill("solid", fgColor="D9D9D9")
BLUE   = PatternFill("solid", fgColor="BDD7EE")
WHITE  = PatternFill("solid", fgColor="FFFFFF")
ORANGE = PatternFill("solid", fgColor="FCE4D6")

HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="2F5496")


def thin_border():
    s = Side(style="thin", color="CCCCCC")
    return Border(left=s, right=s, top=s, bottom=s)


def parse_headline(headline):
    h = (headline or "").strip().lower()
    vehicles = {
        "police":    bool(re.search(r'\bpolice vehicles?\b', h)),
        "fire":      bool(re.search(r'\bfire trucks?\b', h)),
        "ambulance": bool(re.search(r'\bambulances?\b', h)),
        "none":      "no emergency vehicles visible" in h,
    }
    # Singular/plural
    counts = {
        "police":    "police vehicles" in h,   # True=plural, False=singular (if present)
        "fire":      "fire trucks" in h,
        "ambulance": "ambulances" in h,
    }
    is_blocked = "road blocked" in h or "blocked road" in h
    incidents = {
        "crash":        bool(re.search(r'\bcrash\b', h)),
        "pulled_over":  "pulled-over vehicle" in h,
        "blocked_road": is_blocked,
        "construction": "construction" in h,
        "fire_incident":bool(re.search(r'responding to fire\b', h)),
        "crowd":        "crowd" in h,
    }
    return vehicles, counts, incidents


def load_gt(path):
    gt = {}
    with open(path, newline="", encoding="utf-8-sig") as f:
        rows = list(csv.reader(f))
    for idx, row in enumerate(rows):
        if len(row) <= GT_COL_START + len(GT_FIELDS) - 1:
            continue
        rec = {GT_FIELDS[j]: row[GT_COL_START + j].strip().upper() == "TRUE"
               for j in range(len(GT_FIELDS))}
        gt[idx] = rec
    return gt


def gt_summary(rec):
    vehicles = {
        "police":    rec["police_1"] or rec["police_2plus"],
        "fire":      rec["fire_truck_1"] or rec["fire_truck_2plus"],
        "ambulance": rec["ambulance_1"] or rec["ambulance_2plus"],
        "none":      rec["no_emergency_vehicles"],
    }
    gt_plural = {
        "police":    rec["police_2plus"],
        "fire":      rec["fire_truck_2plus"],
        "ambulance": rec["ambulance_2plus"],
    }
    incidents = {
        "crash":        rec["crash"],
        "pulled_over":  rec["pulled_over"],
        "blocked_road": rec["blocked_road"],
        "construction": rec["construction"],
        "fire_incident":rec["fire"],
        "crowd":        rec["crowd"],
    }
    return vehicles, gt_plural, incidents


def load_corrections(path):
    corr = {}
    with open(path, newline="", encoding="utf-8-sig") as f:
        for row in csv.DictReader(f):
            rn = int(row["row"])
            fp_labels = [x.strip() for x in row.get("fp_labels", "").split(";") if x.strip()]
            corr[rn] = {
                "prediction_correct": row.get("prediction_correct", "").strip().lower() == "yes",
                "gt_correct":         row.get("gt_correct", "").strip().lower() == "yes",
                "fp_labels":          fp_labels,
            }
    return corr


def label_list(d):
    return ", ".join(k for k, v in d.items() if v) or "—"


def fetch_image(url, row_num):
    os.makedirs(IMG_CACHE_DIR, exist_ok=True)
    cache_path = os.path.join(IMG_CACHE_DIR, f"row_{row_num}.jpg")
    if os.path.exists(cache_path):
        return cache_path
    try:
        ctx = ssl.create_default_context()
        req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
        with urllib.request.urlopen(req, context=ctx, timeout=15) as resp:
            data = resp.read()
        with open(cache_path, "wb") as f:
            f.write(data)
        return cache_path
    except Exception as e:
        print(f"  row {row_num}: image fetch failed — {e}")
        return None


def build_report():
    gt_all      = load_gt(TESTSHEET_PATH)
    corrections = load_corrections(CORRECTIONS_PATH)

    # Build TP overrides / FN removals
    tp_overrides = set()
    fn_removals  = set()

    pred_by_row = {}
    rows_data = []
    with open(PREDICTIONS_PATH, newline="", encoding="utf-8") as f:
        for row in csv.DictReader(f):
            rows_data.append(row)
            if not row.get("error"):
                pred_by_row[int(row["row"])] = row["headline"]

    for rn, c in corrections.items():
        if not c["prediction_correct"]:
            continue
        for fp_label in c["fp_labels"]:
            scope, label = fp_label.split(":")
            tp_overrides.add((rn, scope, label))
        if rn in gt_all and rn in pred_by_row:
            gt_v, gt_pl, gt_i = gt_summary(gt_all[rn])
            pred_v, pred_c, pred_i = parse_headline(pred_by_row[rn])
            for label, gt_val in gt_v.items():
                if gt_val and not pred_v.get(label):
                    fn_removals.add((rn, "vehicle", label))
            for label, gt_val in gt_i.items():
                if gt_val and not pred_i.get(label):
                    fn_removals.add((rn, "incident", label))

    wb = Workbook()

    # ── Tab 1: Predictions vs GT ─────────────────────────────────────────────
    ws1 = wb.active
    ws1.title = "Predictions vs GT"

    v_labels = ["police", "fire", "ambulance", "none"]
    i_labels = ["crash", "pulled_over", "blocked_road", "construction", "fire_incident", "crowd"]

    headers = (
        ["Row", "Headline",
         "Pred Vehicles", "GT Vehicles",
         "Pred Incidents", "GT Incidents",
         "Overall"] +
        [f"V:{l}" for l in v_labels] +
        [f"I:{l}" for l in i_labels]
    )
    ws1.append(headers)
    for ci, _ in enumerate(headers, 1):
        cell = ws1.cell(1, ci)
        cell.font      = HEADER_FONT
        cell.fill      = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
        cell.border    = thin_border()
    ws1.freeze_panes = "A2"

    vcounts = defaultdict(lambda: [0, 0, 0])
    icounts = defaultdict(lambda: [0, 0, 0])
    # Count accuracy: per vehicle type, among TP rows, how many got singular/plural right
    count_correct = defaultdict(int)   # label -> correct count
    count_total   = defaultdict(int)   # label -> total TPs with count signal

    for data_row in rows_data:
        rn = int(data_row["row"])
        if data_row.get("error") or rn not in gt_all:
            continue

        headline = data_row["headline"]
        pred_v, pred_c, pred_i = parse_headline(headline)
        gt_v, gt_pl, gt_i      = gt_summary(gt_all[rn])

        v_outcomes = {}
        i_outcomes = {}

        for label in v_labels:
            p, g = pred_v.get(label, False), gt_v.get(label, False)
            if (rn, "vehicle", label) in tp_overrides:
                v_outcomes[label] = "TP*"
            elif p and g:
                v_outcomes[label] = "TP"
            elif p and not g:
                v_outcomes[label] = "FP"
            elif not p and g:
                v_outcomes[label] = "TN*" if (rn, "vehicle", label) in fn_removals else "FN"
            else:
                v_outcomes[label] = "TN"

        for label in i_labels:
            p, g = pred_i.get(label, False), gt_i.get(label, False)
            if (rn, "incident", label) in tp_overrides:
                i_outcomes[label] = "TP*"
            elif p and g:
                i_outcomes[label] = "TP"
            elif p and not g:
                i_outcomes[label] = "FP"
            elif not p and g:
                i_outcomes[label] = "TN*" if (rn, "incident", label) in fn_removals else "FN"
            else:
                i_outcomes[label] = "TN"

        # Count accuracy (Option B): only for vehicle TPs
        for label in ("police", "fire", "ambulance"):
            if v_outcomes[label] in ("TP", "TP*"):
                count_total[label] += 1
                pred_plural = pred_c.get(label, False)
                gt_plural   = gt_pl.get(label, False)
                if pred_plural == gt_plural:
                    count_correct[label] += 1

        all_outcomes = list(v_outcomes.values()) + list(i_outcomes.values())
        has_fp = any(o == "FP" for o in all_outcomes)
        has_fn = any(o == "FN" for o in all_outcomes)
        overall = "FP+FN" if (has_fp and has_fn) else "FP" if has_fp else "FN" if has_fn else "✓"

        for label in v_labels:
            o = v_outcomes[label]
            if o in ("TP", "TP*"): vcounts[label][0] += 1
            elif o == "FP":        vcounts[label][1] += 1
            elif o == "FN":        vcounts[label][2] += 1

        for label in i_labels:
            o = i_outcomes[label]
            if o in ("TP", "TP*"): icounts[label][0] += 1
            elif o == "FP":        icounts[label][1] += 1
            elif o == "FN":        icounts[label][2] += 1

        row_vals = (
            [rn, headline,
             label_list(pred_v), label_list(gt_v),
             label_list(pred_i), label_list(gt_i),
             overall] +
            [v_outcomes[l] for l in v_labels] +
            [i_outcomes[l] for l in i_labels]
        )
        ws1.append(row_vals)
        ri = ws1.max_row

        fill_map = {"TP": GREEN, "TP*": GREEN, "FP": RED, "FN": YELLOW, "TN": WHITE, "TN*": WHITE}
        ws1.cell(ri, 7).fill = (GREEN if overall == "✓" else YELLOW if overall in ("FP","FN") else RED)
        ws1.cell(ri, 7).alignment = Alignment(horizontal="center")
        for ci_off, label in enumerate(v_labels):
            c = ws1.cell(ri, 8 + ci_off)
            c.fill = fill_map.get(v_outcomes[label], WHITE)
            c.alignment = Alignment(horizontal="center")
        for ci_off, label in enumerate(i_labels):
            c = ws1.cell(ri, 8 + len(v_labels) + ci_off)
            c.fill = fill_map.get(i_outcomes[label], WHITE)
            c.alignment = Alignment(horizontal="center")
        for ci in range(1, len(headers) + 1):
            ws1.cell(ri, ci).border = thin_border()

    ws1.column_dimensions["A"].width = 6
    ws1.column_dimensions["B"].width = 52
    ws1.column_dimensions["C"].width = 22
    ws1.column_dimensions["D"].width = 22
    ws1.column_dimensions["E"].width = 28
    ws1.column_dimensions["F"].width = 28
    ws1.column_dimensions["G"].width = 9
    for i in range(8, 8 + len(v_labels) + len(i_labels)):
        ws1.column_dimensions[get_column_letter(i)].width = 12

    # ── Tab 2: Performance ───────────────────────────────────────────────────
    ws2 = wb.create_sheet("Performance")
    ws2.column_dimensions["A"].width = 20
    for col in "BCDEFGH":
        ws2.column_dimensions[col].width = 16

    def prf(tp, fp, fn):
        p  = tp / (tp + fp) if (tp + fp) else 0
        r  = tp / (tp + fn) if (tp + fn) else 0
        f1 = 2*p*r / (p + r) if (p + r) else 0
        return round(p, 3), round(r, 3), round(f1, 3)

    def write_table(ws, start_row, title, counts):
        tc = ws.cell(start_row, 1, title)
        tc.font = Font(bold=True, size=12, color="FFFFFF")
        tc.fill = HEADER_FILL
        tc.alignment = Alignment(horizontal="left")
        ws.merge_cells(start_row=start_row, start_column=1, end_row=start_row, end_column=8)
        hr = start_row + 1
        for ci, h in enumerate(["Label","P","R","F1","TP","FP","FN",""], 1):
            c = ws.cell(hr, ci, h)
            c.font = Font(bold=True); c.fill = GREY
            c.alignment = Alignment(horizontal="center"); c.border = thin_border()
        tot_tp = tot_fp = tot_fn = 0
        dr = hr + 1
        for label, (tp, fp, fn) in sorted(counts.items()):
            tot_tp += tp; tot_fp += fp; tot_fn += fn
            p, r, f1 = prf(tp, fp, fn)
            for ci, val in enumerate([label, p, r, f1, tp, fp, fn, ""], 1):
                c = ws.cell(dr, ci, val)
                c.alignment = Alignment(horizontal="center" if ci > 1 else "left")
                c.border = thin_border()
                if ci in (2, 3, 4):
                    c.value = f"{val:.2f}"
                    c.fill = (GREEN if val >= 0.80 else YELLOW if val >= 0.60 else RED)
            dr += 1
        p, r, f1 = prf(tot_tp, tot_fp, tot_fn)
        for ci, val in enumerate(["MICRO AVG", p, r, f1, tot_tp, tot_fp, tot_fn, ""], 1):
            c = ws.cell(dr, ci, val if ci > 1 else val)
            c.font = Font(bold=True); c.fill = BLUE
            c.alignment = Alignment(horizontal="center" if ci > 1 else "left")
            c.border = thin_border()
            if ci in (2, 3, 4):
                c.value = f"{val:.2f}"
        return dr + 2

    r = 1
    ws2.cell(r, 1, "Headline Accuracy Report — 100 records").font = Font(bold=True, size=14)
    r += 1
    ws2.cell(r, 1, "Corrected GT: 29 annotation errors fixed by human review").font = Font(italic=True, color="666666")
    r += 2
    r = write_table(ws2, r, "VEHICLES", vcounts)
    r = write_table(ws2, r, "INCIDENTS", icounts)

    # Overall
    all_tp = sum(v[0] for v in list(vcounts.values()) + list(icounts.values()))
    all_fp = sum(v[1] for v in list(vcounts.values()) + list(icounts.values()))
    all_fn = sum(v[2] for v in list(vcounts.values()) + list(icounts.values()))
    op, orec, of1 = prf(all_tp, all_fp, all_fn)
    ws2.cell(r, 1, "OVERALL").font = Font(bold=True, size=11)
    for ci, (h, v) in enumerate([("P", op), ("R", orec), ("F1", of1),
                                  ("TP", all_tp), ("FP", all_fp), ("FN", all_fn)], 2):
        cell = ws2.cell(r, ci)
        cell.value = f"{v:.2f}" if ci <= 4 else v
        cell.font = Font(bold=True)
        cell.fill = GREEN if (isinstance(v, float) and v >= 0.80) else YELLOW
        cell.alignment = Alignment(horizontal="center"); cell.border = thin_border()
    ws2.cell(r, 1).border = thin_border()
    r += 2

    # Count accuracy section
    ws2.cell(r, 1, "COUNT ACCURACY (singular vs plural — among vehicle TPs only)").font = Font(bold=True, size=12, color="FFFFFF")
    ws2.cell(r, 1).fill = HEADER_FILL
    ws2.merge_cells(start_row=r, start_column=1, end_row=r, end_column=8)
    r += 1
    for ci, h in enumerate(["Vehicle type", "Correct count", "Total TPs", "Accuracy", ""], 1):
        c = ws2.cell(r, ci, h)
        c.font = Font(bold=True); c.fill = GREY
        c.alignment = Alignment(horizontal="center" if ci > 1 else "left")
        c.border = thin_border()
    r += 1
    for label in ("police", "fire", "ambulance"):
        total = count_total[label]
        correct = count_correct[label]
        acc = correct / total if total else 0
        for ci, val in enumerate([label, correct, total, f"{acc:.2f}", ""], 1):
            c = ws2.cell(r, ci, val)
            c.alignment = Alignment(horizontal="center" if ci > 1 else "left")
            c.border = thin_border()
            if ci == 4:
                c.fill = GREEN if acc >= 0.80 else YELLOW if acc >= 0.60 else RED
        r += 1
    total_c = sum(count_correct.values())
    total_t = sum(count_total.values())
    overall_acc = total_c / total_t if total_t else 0
    for ci, val in enumerate(["OVERALL", total_c, total_t, f"{overall_acc:.2f}", ""], 1):
        c = ws2.cell(r, ci, val)
        c.font = Font(bold=True); c.fill = BLUE
        c.alignment = Alignment(horizontal="center" if ci > 1 else "left")
        c.border = thin_border()
        if ci == 4:
            c.fill = GREEN if overall_acc >= 0.80 else YELLOW
    r += 2

    # Legend
    ws2.cell(r, 1, "Legend").font = Font(bold=True); r += 1
    for text, fill in [("TP / TP* (GT-corrected TP)", GREEN),
                       ("FP (predicted but wrong)", RED),
                       ("FN (missed)", YELLOW)]:
        ws2.cell(r, 1, text).fill = fill
        ws2.cell(r, 1).border = thin_border(); r += 1

    # ── Tab 3: Images ────────────────────────────────────────────────────────
    ws3 = wb.create_sheet("Images")
    ws3.column_dimensions["A"].width = 6    # Row #
    ws3.column_dimensions["B"].width = 30   # Image
    ws3.column_dimensions["C"].width = 55   # Headline
    ws3.column_dimensions["D"].width = 20   # Overall

    for ci, h in enumerate(["Row", "Image", "Headline", "Overall"], 1):
        c = ws3.cell(1, ci, h)
        c.font = HEADER_FONT; c.fill = HEADER_FILL
        c.alignment = Alignment(horizontal="center"); c.border = thin_border()
    ws3.freeze_panes = "A2"

    # Pre-build outcomes lookup for colour
    outcomes_by_row = {}
    with open(PREDICTIONS_PATH, newline="", encoding="utf-8") as f:
        for data_row in csv.DictReader(f):
            rn = int(data_row["row"])
            if data_row.get("error") or rn not in gt_all:
                continue
            pred_v, pred_c, pred_i = parse_headline(data_row["headline"])
            gt_v, gt_pl, gt_i      = gt_summary(gt_all[rn])
            has_fp = any(
                (pred_v.get(l) and not gt_v.get(l) and (rn,"vehicle",l) not in tp_overrides)
                for l in v_labels
            ) or any(
                (pred_i.get(l) and not gt_i.get(l) and (rn,"incident",l) not in tp_overrides)
                for l in i_labels
            )
            has_fn = any(
                (not pred_v.get(l) and gt_v.get(l) and (rn,"vehicle",l) not in fn_removals)
                for l in v_labels
            ) or any(
                (not pred_i.get(l) and gt_i.get(l) and (rn,"incident",l) not in fn_removals)
                for l in i_labels
            )
            outcomes_by_row[rn] = "FP+FN" if (has_fp and has_fn) else "FP" if has_fp else "FN" if has_fn else "✓"

    print("Downloading images...")
    IMG_ROW_H = 115   # points (~150px)
    img_ri = 2

    with open(PREDICTIONS_PATH, newline="", encoding="utf-8") as f:
        for idx, data_row in enumerate(csv.DictReader(f), 1):
            rn = int(data_row["row"])
            if data_row.get("error") or rn not in gt_all:
                continue

            print(f"  [{idx}/100] row {rn}", end=" ", flush=True)
            img_path = fetch_image(data_row["url"], rn)

            ws3.row_dimensions[img_ri].height = IMG_ROW_H

            ws3.cell(img_ri, 1, rn).alignment = Alignment(horizontal="center", vertical="center")
            ws3.cell(img_ri, 1).border = thin_border()

            if img_path:
                try:
                    xl_img = XLImage(img_path)
                    xl_img.width  = IMG_W
                    xl_img.height = IMG_H
                    xl_img.anchor = f"B{img_ri}"
                    ws3.add_image(xl_img)
                    print("ok")
                except Exception as e:
                    ws3.cell(img_ri, 2, f"[img error: {e}]")
                    print(f"embed error: {e}")
            else:
                ws3.cell(img_ri, 2, "[unavailable]")
                print("unavailable")

            headline = data_row["headline"]
            hc = ws3.cell(img_ri, 3, headline)
            hc.alignment = Alignment(wrap_text=True, vertical="center")
            hc.border = thin_border()

            overall = outcomes_by_row.get(rn, "")
            oc = ws3.cell(img_ri, 4, overall)
            oc.fill = (GREEN if overall == "✓" else YELLOW if overall in ("FP","FN") else RED if overall == "FP+FN" else WHITE)
            oc.alignment = Alignment(horizontal="center", vertical="center")
            oc.border = thin_border()

            img_ri += 1

    wb.save(OUT_PATH)
    print(f"\nSaved {OUT_PATH}")


if __name__ == "__main__":
    build_report()
